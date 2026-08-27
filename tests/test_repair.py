"""Pandoc-defect repair tests (ticket 06).

Md-side prevention (NormalizeTables / NormalizeImages) as AST transforms;
docx-side repair (FixTable / ClampImageWidths) as XML assertions on the
final DOCX, plus one op-level clamp test with an explicit sectPr (pandoc's
own writer pre-clamps images to its assumed 8400-twip text width, so the
clamp is exercised on a constructed document).
"""

from __future__ import annotations

import struct
import zipfile
import zlib
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from snippet_docx.cli import main
from snippet_docx.ops_docx import ClampImageWidths, ecepdi_docx_ops
from snippet_docx.ops_md_normalize import NormalizeImages, NormalizeTables
from snippet_docx.pandoc_ast import md_to_ast

DATA = Path(__file__).parent / "data"
TABLES_MD = DATA / "tables.md"
SNIPPET = DATA / "snippet.yaml"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# A4 portrait text width (styles fallback: 11906 - 2x1440 twips in EMU)
TEXT_WIDTH_EMU = (11906 - 2880) * 635


def _xml(docx: Path) -> ET.Element:
    with zipfile.ZipFile(docx) as zf:
        return ET.fromstring(zf.read("word/document.xml"))


def _tables(root: ET.Element) -> list[ET.Element]:
    return root.findall(f".//{W}tbl")


def _texts(el: ET.Element) -> list[str]:
    return [t.text or "" for t in el.iter(f"{W}t")]


def _polished(md: Path, snippet: Path, tmp_path: Path,
              monkeypatch: pytest.MonkeyPatch, chdir: Path | None = None) -> Path:
    monkeypatch.chdir(chdir if chdir is not None else DATA)
    draft = tmp_path / "draft.docx"
    final = tmp_path / "final.docx"
    rc = main(["draft", str(md), "--snippet", str(snippet), "--out", str(draft)])
    assert rc == 0
    rc = main(["polish", str(draft), "--snippet", str(snippet), "--out", str(final)])
    assert rc == 0
    return final


def _make_png(path: Path, width: int, height: int) -> None:
    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return (struct.pack(">I", len(data)) + body
                + struct.pack(">I", zlib.crc32(body)))

    raw = b"".join(b"\x00" + b"\x11\x22\x33" * width for _ in range(height))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b""))


def _images(ast: dict) -> list[dict]:
    found: list[dict] = []

    def walk(node: object) -> None:
        if isinstance(node, dict):
            if node.get("t") == "Image":
                found.append(node)
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(ast)
    return found


# ------------------------------------------------------------- md ops

def test_raw_html_table_recovers_to_real_table_ast() -> None:
    ast = md_to_ast("<table>\n<tr><th>机位</th><th>容量</th></tr>\n"
                    "<tr><td>WTG01</td><td>5000</td></tr>\n</table>\n")
    NormalizeTables()(ast, None)
    tables = [b for b in ast["blocks"] if b.get("t") == "Table"]
    assert len(tables) == 1
    cells = [[c[4][0]["c"][0]["c"] if c[4] else ""
              for c in row[1]]
             for tb in tables[0]["c"][4] for row in tb[3]]
    assert cells == [["WTG01", "5000"]]
    head = [[c[4][0]["c"][0]["c"] if c[4] else ""
             for c in row[1]] for row in tables[0]["c"][3][1]]
    assert head == [["机位", "容量"]]


def test_fullwidth_space_simple_table_wreckage_is_rebuilt() -> None:
    # the shape pandoc 3.7 eats: fullwidth-space alignment collapses the
    # whole table into one naked Para (dash row smart-mangled to –/—)
    ast = md_to_ast("机位编号　　装机容量　　年上网电量\n"
                    "--------　　----------　　----------\n"
                    "WTG01　　　 5000　　　　 1234\n")
    kinds = [b["t"] for b in ast["blocks"]]
    assert kinds == ["Para"]  # the defect: no Table at all
    NormalizeTables()(ast, None)
    tables = [b for b in ast["blocks"] if b.get("t") == "Table"]
    assert len(tables) == 1
    head = [[ "".join(e.get("c", "") for e in c[4][0]["c"])
              for c in row[1]] for row in tables[0]["c"][3][1]]
    body = [["".join(e.get("c", "") for e in c[4][0]["c"])
             for c in row[1]] for tb in tables[0]["c"][4] for row in tb[3]]
    assert head == [["机位编号", "装机容量", "年上网电量"]]
    assert body == [["WTG01", "5000", "1234"]]


def test_wellformed_simple_table_is_a_verified_noop() -> None:
    md = "机位      容量\n----      ----\nT01       7.7\nT02       7.7\n"
    ast = md_to_ast(md)
    before = repr(ast)
    NormalizeTables()(ast, None)
    assert repr(ast) == before


def test_prose_paragraph_is_never_eaten() -> None:
    md = "风速数据 如下所示。\n\n湍流强度 0.12 与 平均风速 6.8 相关。\n"
    ast = md_to_ast(md)
    before = repr(ast)
    NormalizeTables()(ast, None)
    assert repr(ast) == before


def test_normalizeimages_injects_real_pixel_size(tmp_path: Path,
                                                 monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    _make_png(tmp_path / "wide.png", 2000, 500)
    ast = md_to_ast("![wide](wide.png)\n")
    NormalizeImages()(ast, None)
    assert _images(ast)[0]["c"][0][2] == [["width", "2000px"], ["height", "500px"]]


def test_normalizeimages_default_when_file_unresolvable(tmp_path: Path,
                                                        monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    ast = md_to_ast("![missing](nope.png)\n")
    NormalizeImages()(ast, None)
    # width-only default: pandoc preserves the aspect instead of guessing a
    # height for a file it cannot read
    assert _images(ast)[0]["c"][0][2] == [["width", "600px"]]


def test_normalizeimages_leaves_sized_images_alone(tmp_path: Path,
                                                   monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    _make_png(tmp_path / "wide.png", 2000, 500)
    ast = md_to_ast("![sized](wide.png){width=3cm}\n")
    NormalizeImages()(ast, None)
    assert _images(ast)[0]["c"][0][2] == [["width", "3cm"]]


# --------------------------------------------------------- op wiring

def test_strategy_normalize_seam_is_wired() -> None:
    from snippet_docx.strategies.ecepdi import EcepdiStrategy
    ops = EcepdiStrategy().md_normalize_ops()
    assert [type(op).__name__ for op in ops] == ["NormalizeTables", "NormalizeImages"]


def test_docx_ops_repair_before_strip() -> None:
    names = [type(op).__name__ for op in ecepdi_docx_ops()]
    assert names.index("FixTable") < names.index("StripAnchors")
    assert names.index("ClampImageWidths") < names.index("StripAnchors")
    assert names[-1] == "StripAnchors"


# ------------------------------------------------------------- e2e

def test_simple_tables_survive_into_docx_as_real_tables(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    final = _polished(TABLES_MD, SNIPPET, tmp_path, monkeypatch)
    tables = _tables(_xml(final))
    assert len(tables) == 5  # front-matter, pipe, ascii-simple, wreckage, html
    all_cells = [_texts(t) for t in tables]
    # well-formed simple table (verified no-op shape): real w:tbl, right cells
    assert ["机位", "容量", "T01", "7.7", "T02", "7.7"] in all_cells
    # fullwidth-space wreckage recovered with exact cells
    assert ["机位编号", "装机容量", "年上网电量",
            "WTG01", "5000", "1234", "WTG02", "5000", "2345"] in all_cells
    # raw HTML table recovered (pandoc's docx writer drops it without the op)
    assert ["机位", "容量", "WTG01", "5000"] in all_cells


def test_content_table_gets_borders_centering_and_content_widths(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    final = _polished(TABLES_MD, SNIPPET, tmp_path, monkeypatch)
    tables = _tables(_xml(final))
    pipe = next(t for t in tables if _texts(t)[0] == "机位编号"
                and len(_texts(t)) > 10)
    tblpr = pipe.find(f"{W}tblPr")
    borders = tblpr.find(f"{W}tblBorders")
    assert borders is not None
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(f"{W}{edge}")
        assert el is not None, f"missing border edge {edge}"
        assert el.get(f"{W}val") == "single"
        assert el.get(f"{W}sz") == "4"
        assert el.get(f"{W}color") == "000000"
    jc = tblpr.find(f"{W}jc")
    assert jc is not None and jc.get(f"{W}val") == "center"
    # content-derived grid, twip rules at 9 pt (+108 padding; floors):
    #   机位编号/WTG02号机 -> 5ascii+2cjk=400+320+108          = 828
    #   numeric col (7.7) floor                                 = 1200
    #   年上网电量 (GWh) -> 5cjk+6ascii=800+480+108             = 1388
    #   备注/正常 -> 2cjk+108=428 -> text floor                 = 500
    grid = [int(c.get(f"{W}w")) for c in pipe.find(f"{W}tblGrid")]
    assert grid == [828, 1200, 1388, 500]
    widths = [int(tc.find(f"{W}tcPr/{W}tcW").get(f"{W}w"))
              for tc in pipe.findall(f"{W}tr")[0].findall(f"{W}tc")]
    assert widths == grid


def test_cell_text_carries_configured_9pt_style(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    final = _polished(TABLES_MD, SNIPPET, tmp_path, monkeypatch)
    tables = _tables(_xml(final))
    pipe = next(t for t in tables if _texts(t)[0] == "机位编号"
                and len(_texts(t)) > 10)
    # every cell paragraph points at the named `table` style ...
    paras = list(pipe.iter(f"{W}p"))
    assert paras
    for para in paras:
        pstyle = para.find(f"{W}pPr/{W}pStyle")
        assert pstyle is not None and pstyle.get(f"{W}val") == "table"
        # ... and no direct run formatting fights it (styles engine contract)
        for run in para.findall(f"{W}r"):
            rpr = run.find(f"{W}rPr")
            if rpr is None:
                continue
            assert rpr.find(f"{W}sz") is None
            rfonts = rpr.find(f"{W}rFonts")
            if rfonts is not None:
                assert not [a for a in rfonts.attrib if a != f"{W}hint"]
    # ... and the named style itself is the configured 9pt 宋体/TNR
    with zipfile.ZipFile(final) as zf:
        styles_xml = ET.fromstring(zf.read("word/styles.xml"))
    table_style = next(s for s in styles_xml.findall(f"{W}style")
                       if s.get(f"{W}styleId") == "table")
    fonts = table_style.find(f"{W}rPr/{W}rFonts")
    assert fonts.get(f"{W}ascii") == "Times New Roman"
    assert fonts.get(f"{W}eastAsia") == "宋体"
    assert table_style.find(f"{W}rPr/{W}sz").get(f"{W}val") == "18"  # 9 pt


def test_front_matter_layout_table_is_not_mangled(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    final = _polished(TABLES_MD, SNIPPET, tmp_path, monkeypatch)
    tables = _tables(_xml(final))
    fm = next(t for t in tables if _texts(t)[0] == "项目名称")
    tblpr = fm.find(f"{W}tblPr")
    assert tblpr.find(f"{W}tblBorders") is None   # no borders forced
    assert tblpr.find(f"{W}jc") is None           # not centered
    # pandoc's own equal grid untouched (no content-width rewrite)
    grid = [int(c.get(f"{W}w")) for c in fm.find(f"{W}tblGrid")]
    assert grid == [3960, 3960]
    assert _texts(fm) == ["项目名称", "某风电场微观选址报告",
                          "建设地点", "新疆某地", "编制", "设计院"]


def test_oversized_image_clamped_to_text_width(
        tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    _make_png(tmp_path / "huge.png", 2400, 600)
    md = tmp_path / "img.md"
    md.write_text("![wide](huge.png)\n", encoding="utf-8")
    emitted = tmp_path / "anchored.md"
    draft = tmp_path / "draft.docx"
    final = tmp_path / "final.docx"
    rc = main(["draft", str(md), "--snippet", str(SNIPPET),
               "--out", str(draft), "--emit-md", str(emitted)])
    assert rc == 0
    # NormalizeImages recorded the real pixel size on the emitted markdown
    # (pandoc's writer quotes attr values: width="2400px" height="600px")
    emitted_md = emitted.read_text(encoding="utf-8")
    assert "width=" in emitted_md and "2400px" in emitted_md
    assert "height=" in emitted_md and "600px" in emitted_md
    rc = main(["polish", str(draft), "--snippet", str(SNIPPET), "--out", str(final)])
    assert rc == 0
    root = _xml(final)
    for extent in root.iter(f"{WP}extent"):
        assert int(extent.get("cx")) <= TEXT_WIDTH_EMU
        assert int(extent.get("cy")) <= TEXT_WIDTH_EMU
    # every DrawingML transform agrees with its extent
    for ext in root.iter(f"{A}ext"):
        assert int(ext.get("cx")) <= TEXT_WIDTH_EMU


# --------------------------------------------------- op-level clamp

def test_clamp_image_widths_clamps_to_section_text_width(tmp_path: Path) -> None:
    from docx import Document
    from docx.shared import Emu, Twips
    _make_png(tmp_path / "big.png", 1600, 400)
    doc = Document()
    sec = doc.sections[-1]
    sec.page_width = Twips(11906)      # A4 portrait, explicit geometry
    sec.page_height = Twips(16838)
    sec.left_margin = Twips(1440)
    sec.right_margin = Twips(1440)
    para = doc.add_paragraph()
    para.add_run().add_picture(str(tmp_path / "big.png"),
                               width=Emu(9144000), height=Emu(2286000))
    doc.save(str(tmp_path / "in.docx"))

    doc = Document(str(tmp_path / "in.docx"))
    ClampImageWidths()(doc, None)
    shape = doc.inline_shapes[0]
    assert shape.width == TEXT_WIDTH_EMU
    expected_cy = round(2286000 * TEXT_WIDTH_EMU / 9144000)
    assert abs(shape.height - expected_cy) <= 1  # height scales proportionally
    extents = [extent.attrib for extent in doc.element.body.iter(f"{WP}extent")]
    transforms = [ext.attrib for ext in doc.element.body.iter(f"{A}ext")]
    assert all(int(e["cx"]) == TEXT_WIDTH_EMU for e in extents + transforms)


def test_clamp_leaves_fitting_images_alone(tmp_path: Path) -> None:
    from docx import Document
    from docx.shared import Emu, Twips
    _make_png(tmp_path / "small.png", 400, 100)
    doc = Document()
    sec = doc.sections[-1]
    sec.page_width = Twips(11906)
    sec.left_margin = Twips(1440)
    sec.right_margin = Twips(1440)
    doc.add_paragraph().add_run().add_picture(
        str(tmp_path / "small.png"), width=Emu(2000000))
    doc.save(str(tmp_path / "fit.docx"))
    doc = Document(str(tmp_path / "fit.docx"))
    ClampImageWidths()(doc, None)
    assert doc.inline_shapes[0].width == 2000000  # untouched
